"""Compute MTD variances as a P&L: trial balance actuals vs the budget workbook.

Usage: variance.py tb.json budget.xlsx <period> <flag_threshold_k> outdir/ \
                   "<entity>" "<footer>" pnl.json [commentary.md]

pnl.json groups every trial-balance line (written by run_step2.py from
close.toml [pnl]): revenue, cost_of_sales, operating_expenses, below_ebitda.
Each line sits in exactly one group. The report builds the standard
subtotals from those groups: total revenue, gross margin, total operating
expenses, EBITDA and net result.

Variance = actual - budget on every row. F/A says what it does to profit: a
revenue or profit row above budget is F (favourable); a cost row above
budget is A (adverse). Lines are flagged at the materiality threshold.

Writes outdir/variance.json and "outdir/Variance Report - <period>.xlsx",
all in thousands. With a commentary file, the workbook gains a Commentary
sheet holding it verbatim: the reviewer signs the numbers and the words in
one document. Requires openpyxl.
"""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

GROUPS = ("revenue", "cost_of_sales", "operating_expenses", "below_ebitda")
INCOME = {"revenue"}


def read_budget(path: str, period: str) -> dict:
    ws = openpyxl.load_workbook(path).active
    header = None
    budgets = {}
    for row in ws.iter_rows(values_only=True):
        if row[0] == "Line":
            header = list(row)
            continue
        if header and row[0] and isinstance(row[0], str):
            month_idx = header.index(period)
            if row[month_idx] is not None:
                budgets[row[0]] = float(row[month_idx])
    if not budgets:
        raise SystemExit(f"no budget column for {period} in {path}")
    return budgets


def effect(variance: float, income: bool) -> str:
    """F or A for profit: income above budget is good, cost above budget is not."""
    if variance == 0:
        return ""
    return "F" if (variance > 0) == income else "A"


def subtotals(lines: list[dict]) -> list[dict]:
    """The standard P&L subtotals, each summed from the line rows only."""
    def total(group: str, key: str) -> float:
        return sum(l[key] for l in lines if l["group"] == group)

    rows = []
    for key in ("mtd_actual_k", "mtd_budget_k"):
        rev, cos = total("revenue", key), total("cost_of_sales", key)
        opex, below = total("operating_expenses", key), total("below_ebitda", key)
        gm = rev - cos
        ebitda = gm - opex
        rows.append({"Total revenue": rev, "Gross margin": gm,
                     "Total operating expenses": opex, "EBITDA": ebitda,
                     "Net result": ebitda - below})
    actual, budget = rows
    out = []
    for name in actual:
        var = round(actual[name] - budget[name], 1)
        out.append({"name": name, "mtd_actual_k": round(actual[name], 1),
                    "mtd_budget_k": round(budget[name], 1), "variance_k": var,
                    "effect": effect(var, income=name != "Total operating expenses")})
    return out


def layout(lines: list[dict], subs: dict) -> list[tuple[str, dict]]:
    """Rows in reading order: each group's lines, then the subtotal it feeds."""
    after = {"revenue": ["Total revenue"], "cost_of_sales": ["Gross margin"],
             "operating_expenses": ["Total operating expenses", "EBITDA"],
             "below_ebitda": ["Net result"]}
    rows = []
    for group in GROUPS:
        rows += [("line", l) for l in lines if l["group"] == group]
        rows += [("subtotal", subs[name]) for name in after[group]]
    return rows


def main(tb_path: str, budget_path: str, period: str, threshold_k: str, outdir: str,
         entity: str, footer: str, pnl_path: str, commentary_path: str = "") -> None:
    tb = json.load(open(tb_path, encoding="utf-8"))
    assert tb["period"] == period, f"trial balance is {tb['period']}, wanted {period}"
    budgets = read_budget(budget_path, period)
    pnl = json.load(open(pnl_path, encoding="utf-8"))
    group_of = {line: g for g in GROUPS for line in pnl.get(g, [])}
    threshold = float(threshold_k)

    lines = []
    for line, vals in tb["lines"].items():
        if line not in budgets:
            raise SystemExit(f"line missing from budget: {line}")
        if line not in group_of:
            raise SystemExit(f"line not placed in the P&L (close.toml [pnl]): {line}")
        actual, budget = float(vals["mtd_actual"]), budgets[line]
        var = round(actual - budget, 1)
        pct = round(100 * var / budget, 1) if budget else None
        group = group_of[line]
        lines.append({"line": line, "group": group,
                      "mtd_actual_k": actual, "mtd_budget_k": budget,
                      "variance_k": var, "variance_pct": pct,
                      "effect": effect(var, income=group in INCOME),
                      "flagged": abs(var) >= threshold})
    subs = subtotals(lines)
    by_name = {s["name"]: s for s in subs}
    result = {
        "report": "MTD variance vs budget, as a P&L", "period": period,
        "units": "GBP thousands", "flag_threshold_k": threshold,
        "line_count": len(lines),
        "flagged_count": sum(1 for l in lines if l["flagged"]),
        "ebitda_actual_k": by_name["EBITDA"]["mtd_actual_k"],
        "ebitda_variance_k": by_name["EBITDA"]["variance_k"],
        "net_result_variance_k": by_name["Net result"]["variance_k"],
        "lines": lines,
        "subtotals": subs,
    }
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    json.dump(result, open(out / "variance.json", "w", encoding="utf-8", newline="\n"), indent=1)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Variance"
    ws["A1"] = f"{entity} — Variance Report {period} (thousands, MTD)"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = (f"Actual per trial balance; budget per {Path(budget_path).stem}. Variance = actual "
                f"- budget; F/A = effect on profit. Flag threshold ±{threshold:,.0f}. {footer}")
    ws.append([])
    ws.append(["Line", "Actual", "Budget", "Variance", "Var %", "F/A", "Flag"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for kind, r in layout(lines, by_name):
        if kind == "line":
            ws.append([r["line"], r["mtd_actual_k"], r["mtd_budget_k"], r["variance_k"],
                       r["variance_pct"], r["effect"], "FLAG" if r["flagged"] else ""])
        else:
            ws.append([r["name"], r["mtd_actual_k"], r["mtd_budget_k"], r["variance_k"],
                       None, r["effect"], None])
            for c in ws[ws.max_row]:
                c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 38
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 11
    for row in ws.iter_rows(min_row=5, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = "#,##0"
    if commentary_path:
        notes = wb.create_sheet("Commentary")
        notes.column_dimensions["A"].width = 110
        for para in Path(commentary_path).read_text(encoding="utf-8").strip().splitlines():
            notes.append([para])
    wb.save(out / f"Variance Report - {period}.xlsx")
    e = by_name["EBITDA"]
    print(f"{result['flagged_count']}/{result['line_count']} lines flagged at ±{threshold:,.0f}k; "
          f"EBITDA {e['mtd_actual_k']:,.0f} vs {e['mtd_budget_k']:,.0f}, "
          f"{e['variance_k']:+,.1f}k {e['effect']} -> {out}")


if __name__ == "__main__":
    main(*sys.argv[1:10])
