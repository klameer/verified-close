"""The close, end to end, and every way it is meant to refuse."""
import hashlib
import json

from openpyxl import load_workbook


def sha256(path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_golden_set_is_green(close):
    out = close.ok("run_goldens.py")
    assert "golden set: 8/8 green" in out


def test_figures_are_byte_identical_to_the_recorded_run(close):
    """Same inputs, same bytes, on any machine: the sealed figures of showcase/may-26."""
    close.ok("run_step1.py")
    close.ok("sign_step1.py", "KL", "K. Lameer", "reissue")
    close.ok("run_step2.py")
    assert sha256(close.work / "step1" / "reconciliation.json") == \
        "44fc053f059d5f0f36be4df54d9a03b9026645837c6473469c569bf23c570dcc"
    assert sha256(close.work / "step2" / "variance.json") == \
        "9e007c8f0143ba771165547281532b244034e250319cf83ccf256e15d90fccc8"


def test_full_close_end_to_end(close):
    close.sign_step1()
    close.sign_step2()
    close.sign_step3()
    out = close.ok("close_out.py")
    assert "seals verified" in out
    records = [json.loads(l) for l in (close.work / "audit" / "close-log.jsonl").read_text().splitlines()]
    assert [r["step"] for r in records] == ["step1", "step2", "step3", "closeout"]
    assert records[0]["note"] == "cheque 004182: stop and reissue"
    for name in ["Bank Reconciliation - Operating Account - May-26 SIGNED KL.xlsx",
                 "Variance Report - May-26 SIGNED KL.xlsx",
                 "Management Pack - May-26 SIGNED KL.pptx",
                 "Close Binder - May-26.pdf", "close-log.jsonl"]:
        assert (close.filing / name).exists(), name
    board = close.ok("status.py")
    assert "[x] Close-out (audit binder)" in board and "ready" not in board
    assert "already closed out" in close.run("close_out.py").stdout


def test_signed_rec_feeds_next_month(close):
    close.sign_step1()
    ws = load_workbook(close.filing / "Bank Reconciliation - Operating Account - May-26 SIGNED KL.xlsx").active
    text = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("Reviewed: K. Lameer" in t for t in text)
    assert not any("DRAFT" in t for t in text)


# ---- the gates -------------------------------------------------------------

def test_step2_refuses_until_step1_signed(close):
    close.ok("run_step1.py")
    r = close.run("run_step2.py")
    assert r.returncode == 1 and "GATE" in r.stdout


def test_step3_refuses_until_step2_signed(close):
    close.sign_step1()
    close.write("step3/commentary.txt", "Cash reconciled.\n")
    r = close.run("run_step3.py")
    assert r.returncode == 1 and "GATE" in r.stdout and "Step 2" in r.stdout


def test_close_out_refuses_until_all_signed(close):
    close.sign_step1()
    r = close.run("close_out.py")
    assert r.returncode == 1 and "GATE" in r.stdout


def test_stale_item_cannot_be_signed_without_a_decision(close):
    close.ok("run_step1.py")
    r = close.run("sign_step1.py", "KL", "K. Lameer")
    assert r.returncode == 1 and "decision required" in r.stdout and "004182" in r.stdout


def test_a_signed_step_never_reruns(close):
    close.sign_step1()
    r = close.run("run_step1.py")
    assert r.returncode == 1 and "already signed" in r.stdout


def test_pack_needs_its_commentary_first(close):
    close.sign_step1()
    close.sign_step2()
    r = close.run("run_step3.py")
    assert r.returncode == 1 and "commentary" in r.stdout


# ---- the seals and the checks ---------------------------------------------

def test_editing_a_signed_file_breaks_the_seal(close):
    close.sign_step1()
    signed = close.filing / "Bank Reconciliation - Operating Account - May-26 SIGNED KL.xlsx"
    wb = load_workbook(signed)
    wb.active["B5"] = 1.0
    wb.save(signed)
    r = close.run("run_step2.py")
    assert r.returncode == 1 and "SEAL BROKEN" in r.stdout


def test_editing_signed_figures_breaks_the_seal(close):
    close.sign_step1()
    close.sign_step2()
    var = close.work / "step2" / "variance.json"
    var.write_text(var.read_text().replace('"total_variance_k": 545.0', '"total_variance_k": 445.0'))
    close.write("step3/commentary.txt", "Cash reconciled.\n")
    r = close.run("run_step3.py")
    assert r.returncode == 1 and "SEAL BROKEN" in r.stdout


def test_figures_edited_before_signoff_fail_the_recheck(close):
    close.ok("run_step1.py")
    rec = close.work / "step1" / "reconciliation.json"
    data = json.loads(rec.read_text())
    data["reconciliation_statement"]["less_unpresented_payments_gbp"] += 100
    rec.write_text(json.dumps(data))
    r = close.run("sign_step1.py", "KL", "K. Lameer", "reissue")
    assert r.returncode == 1 and "CHECKS FAILED" in r.stdout


def test_a_broken_bank_statement_stops_step1(close):
    export = close.path / "02 System Exports" / "May-26" / "Bank Statement - May-26.json"
    data = json.loads(export.read_text())
    data["rows"][10]["balance_gbp"] = str(float(data["rows"][10]["balance_gbp"]) + 50)
    export.write_text(json.dumps(data))
    r = close.run("run_step1.py")
    assert r.returncode == 1 and "C1 running balance breaks" in r.stdout
    assert not any("DRAFT" in p.name for p in close.filing.iterdir())


def test_materiality_change_keeps_goldens_green(close):
    toml = close.path / "close.toml"
    toml.write_text(toml.read_text(encoding="utf-8").replace("materiality_k = 100", "materiality_k = 50"),
                    encoding="utf-8")
    assert "8/8 green" in close.ok("run_goldens.py")


# ---- reset -----------------------------------------------------------------

def test_reset_removes_only_run_outputs(close):
    close.sign_step1()
    estate = {p.name for p in close.filing.iterdir()} - {
        "Bank Reconciliation - Operating Account - May-26 DRAFT.xlsx",
        "Bank Reconciliation - Operating Account - May-26 SIGNED KL.xlsx"}
    assert "dry run" in close.ok("reset_close.py")
    close.ok("reset_close.py", "--yes")
    assert {p.name for p in close.filing.iterdir()} == estate
    assert not close.work.exists()
    assert (close.path / "05 Month-End" / "Apr-26").exists()
